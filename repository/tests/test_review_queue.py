import copy
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from PIL import Image
import tifffile

from app.project import Project
from app.review_queue import box_signature, mask_signature, segment_approved, export_approved


class SamEngine:
    def __init__(self, stop=None, fail=False):
        self.config = {'resolved_paths': {'sam_model': 'sam.pt'}, 'inference': {'device': 'cpu'}}
        self.stop, self.fail = stop, fail
        self.calls, self.released = [], False

    def segment(self, image, boxes, scores, progress):
        self.calls.append((image.copy(), boxes.copy(), scores.copy()))
        if self.fail and len(self.calls) == 1:
            raise RuntimeError('inference failed')
        labels = np.zeros(image.shape[:2], dtype=np.uint32)
        labels[2:5, 2:5] = 3
        if self.stop is not None:
            self.stop.set()
        return labels, [{'label_id': 3}]

    def release_sam(self):
        self.released = True


class ReviewQueueTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        images = []
        for i in range(3):
            image = self.root / f'{i}.png'
            Image.fromarray(np.full((16, 20), i + 10, dtype=np.uint8)).save(image)
            images.append(image)
        self.project = Project.create(self.root / 'project', images)
        self.prompts = {'boxes_xyxy': [[2., 2., 8., 9.]], 'confidence': [.73],
                        'source': ['manual'], 'pending_delete': [False]}
        for i in range(3):
            self.project.save(i, {'status': 'ready_sam', 'boxes': copy.deepcopy(self.prompts),
                                  'box_approval': box_signature(self.prompts), 'corrected': True, 'radius': 0,
                                  'pixel_size': .5 + i, 'unit': 'um'})
        self.stop = threading.Event()

    def tearDown(self):
        self.tmp.cleanup()

    def approve_exports(self):
        segment_approved(self.project, SamEngine(), self.stop, lambda event: None)
        for i in range(3):
            metadata, labels = self.project.load(i)
            metadata.update(status='ready_export', mask_approval=mask_signature(labels, metadata['pixel_size'], metadata['unit']))
            self.project.save(i, metadata, labels)

    def test_sam_stop_resume_and_same_prepared_pixels(self):
        engine = SamEngine(stop=self.stop)
        result = segment_approved(self.project, engine, self.stop, lambda event: None)
        self.assertEqual(result['completed'], 1)
        self.assertTrue(engine.released)
        self.assertEqual(self.project.statuses(), ['segmented', 'ready_sam', 'ready_sam'])
        np.testing.assert_array_equal(engine.calls[0][0], np.full((16, 20), 10 / 255, dtype=np.float32))
        np.testing.assert_array_equal(engine.calls[0][2], [1.])
        self.stop.clear()
        engine = SamEngine()
        result = segment_approved(Project(self.project.path), engine, self.stop, lambda event: None)
        self.assertEqual(result['completed'], 2)
        self.assertEqual(len(engine.calls), 2)

    def test_unapproved_skipped_and_failed_retry(self):
        metadata, labels = self.project.load(2)
        metadata['status'] = 'detected'
        self.project.save(2, metadata)
        result = segment_approved(self.project, SamEngine(fail=True), self.stop, lambda event: None)
        self.assertEqual((result['completed'], result['failed']), (1, 1))
        self.assertEqual(self.project.statuses(), ['ready_sam', 'segmented', 'detected'])
        self.assertIn('queue_error', self.project.load(0)[0])
        result = segment_approved(self.project, SamEngine(), self.stop, lambda event: None)
        self.assertEqual(result['completed'], 1)
        self.assertNotIn('queue_error', self.project.load(0)[0])

    def test_stale_boxes_rejected(self):
        metadata, _ = self.project.load(0)
        metadata['boxes']['boxes_xyxy'][0][0] = 3
        self.project.save(0, metadata)
        result = segment_approved(self.project, SamEngine(), self.stop, lambda event: None)
        self.assertEqual(result['failed'], 1)
        self.assertIn('changed after approval', self.project.load(0)[0]['queue_error']['message'])

    def test_checkpoint_failure_does_not_destroy_approval(self):
        with patch.object(self.project, 'save', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                segment_approved(self.project, SamEngine(), self.stop, lambda event: None)
        self.assertEqual(self.project.statuses(), ['ready_sam'] * 3)

    def test_export_skip_replace_and_calibration(self):
        self.approve_exports()
        destination = self.root / 'exports'
        existing = destination / '0'
        existing.mkdir(parents=True)
        (existing / 'notes.txt').write_text('Keep this')
        (existing / 'labels.tif').write_bytes(b'old result')
        result = export_approved(self.project, destination, ['labels', 'csv'], False, self.stop, lambda e: None)
        self.assertEqual((result['completed'], result['skipped']), (2, 1))
        self.assertEqual(self.project.statuses(), ['ready_export', 'complete', 'complete'])
        self.assertEqual((existing / 'labels.tif').read_bytes(), b'old result')
        result = export_approved(self.project, destination, ['labels', 'csv'], True, self.stop, lambda e: None)
        self.assertEqual(result['completed'], 1)
        self.assertEqual((existing / 'notes.txt').read_text(), 'Keep this')
        np.testing.assert_array_equal(tifffile.imread(existing / 'labels.tif'), self.project.load(0)[1])
        import json
        self.assertEqual(json.loads((destination / '1' / 'session.json').read_text())['pixel_size'], 1.5)

    def test_changed_mask_and_calibration_rejected(self):
        self.approve_exports()
        metadata, labels = self.project.load(0)
        labels[0, 0] = 10
        self.project.save(0, metadata, labels)
        metadata, labels = self.project.load(1)
        metadata['pixel_size'] = 5.
        self.project.save(1, metadata, labels)
        result = export_approved(self.project, self.root / 'exports', ['labels'], False, self.stop, lambda e: None)
        self.assertEqual((result['completed'], result['failed']), (1, 2))

    def test_export_name_collisions_and_project_protection(self):
        self.approve_exports()
        for i, name in enumerate(['same', 'SAME', 'images']):
            metadata, labels = self.project.load(i)
            metadata['export_name'] = name
            self.project.save(i, metadata, labels)
        result = export_approved(self.project, self.project.root, ['labels'], True, self.stop, lambda e: None)
        self.assertEqual(result['failed'], 3)
        self.assertFalse((self.project.root / 'same').exists())
        self.project.read_image(0)

    def test_export_stop_keeps_remaining_approved(self):
        self.approve_exports()
        def progress(event):
            if event['stage'] == 'Saved':
                self.stop.set()
        result = export_approved(self.project, self.root / 'exports', ['labels'], False, self.stop, progress)
        self.assertEqual(result['completed'], 1)
        self.assertTrue(result['stopped'])
        self.assertEqual(self.project.statuses(), ['complete', 'ready_export', 'ready_export'])

    def test_export_all_merge_reexport_and_legacy_names(self):
        import csv
        from openpyxl import load_workbook
        self.approve_exports()
        metadata, labels = self.project.load(0)
        metadata['export_name'] = '0001_0'
        self.project.save(0, metadata, labels)
        root = self.root / 'combined'
        result = export_approved(self.project, root, ['labels'], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertEqual(result['completed'], 3)
        self.assertIn('Combined CSV and Excel saved', result['merge_note'])
        with (root / 'all_measurements.csv').open(newline='') as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual([r['image_name'] for r in rows], ['0.png', '1.png', '2.png'])
        self.assertEqual([float(r['area']) for r in rows], [2.25, 20.25, 56.25])
        for i in range(3):
            self.assertTrue((root / str(i) / 'measurements.csv').is_file())
            self.assertTrue((root / str(i) / 'measurements.xlsx').is_file())
        workbook = load_workbook(root / 'all_measurements.xlsx', data_only=True)
        self.assertEqual(workbook.active.cell(2, 8).value, 2.25)
        self.assertEqual(workbook.active.cell(2, 1).value, '0.png')
        workbook.close()
        result = export_approved(self.project, self.root / 'second', ['csv'], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertEqual(result['completed'], 3)
        # Skipped image tables can be merged without re-exporting their masks.
        (root / 'all_measurements.csv').unlink()
        (root / 'all_measurements.xlsx').unlink()
        result = export_approved(self.project, root, [], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertEqual(result['skipped'], 3)
        self.assertTrue((root / 'all_measurements.csv').is_file())

    def test_all_requires_review_and_exports_empty_images(self):
        with self.assertRaisesRegex(ValueError, 'Review all images'):
            export_approved(self.project, self.root / 'out', ['csv'], False, self.stop, lambda e: None, all_images=True)
        self.approve_exports()
        metadata, _ = self.project.load(1)
        metadata['status'] = 'empty'
        self.project.save(1, metadata)
        result = export_approved(self.project, self.root / 'out', ['labels'], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertEqual(result['completed'], 3)
        self.assertFalse(tifffile.imread(self.root / 'out' / '1' / 'labels.tif').any())
        self.assertEqual(self.project.load(1)[0]['status'], 'empty')
        self.assertIsNone(self.project.load(1)[1])

    def test_combined_not_published_after_failure_or_stop(self):
        self.approve_exports()
        metadata, labels = self.project.load(0)
        labels[0, 0] = 99
        self.project.save(0, metadata, labels)
        result = export_approved(self.project, self.root / 'out', ['csv'], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertEqual(result['failed'], 1)
        self.assertFalse((self.root / 'out' / 'all_measurements.csv').exists())
        self.stop.set()
        result = export_approved(self.project, self.root / 'stopped', ['csv'], False, self.stop, lambda e: None,
                                 all_images=True, merge=True)
        self.assertTrue(result['stopped'])
        self.assertFalse((self.root / 'stopped' / 'all_measurements.xlsx').exists())

    def test_combined_writer_failure_preserves_individual_and_previous_files(self):
        from app.core import write_combined_measurements
        self.approve_exports()
        root = self.root / 'out'
        export_approved(self.project, root, ['csv'], False, self.stop, lambda e: None, all_images=True, merge=True)
        before = {p.name: p.read_bytes() for p in root.glob('all_measurements.*')}
        with patch('app.core.write_measurement_xlsx', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                write_combined_measurements(root, ['image_name'], [['example']], True)
        self.assertEqual(before, {p.name: p.read_bytes() for p in root.glob('all_measurements.*')})
        self.assertTrue((root / '0' / 'measurements.xlsx').is_file())

    def test_progress_reads_do_not_overlap_checkpoint_replacement(self):
        from concurrent.futures import ThreadPoolExecutor
        opened, release, replacing = threading.Event(), threading.Event(), threading.Event()
        original_load = np.load
        import os
        original_replace = os.replace
        def blocked_load(*args, **kwargs):
            archive = original_load(*args, **kwargs)
            opened.set()
            if not release.wait(5):
                archive.close()
                raise RuntimeError('Test timed out')
            return archive
        def replacing_file(*args):
            replacing.set()
            return original_replace(*args)
        metadata, labels = self.project.load(0)
        with ThreadPoolExecutor(max_workers=2) as pool:
            with patch('app.project.np.load', side_effect=blocked_load), \
                 patch('app.project.os.replace', side_effect=replacing_file):
                reader = pool.submit(self.project.summaries)
                self.assertTrue(opened.wait(5))
                writer = pool.submit(self.project.save, 0, metadata, labels)
                try:
                    self.assertFalse(replacing.wait(.1))
                finally:
                    release.set()
                reader.result(timeout=5)
                writer.result(timeout=5)
                self.assertTrue(replacing.is_set())


if __name__ == '__main__':
    unittest.main()
