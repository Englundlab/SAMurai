"""Inference and export, independent of the napari interface."""
from __future__ import annotations

import csv
import gc
import json
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import tifffile
from PIL import Image
from scipy.special import ive
from skimage.measure import regionprops
from .display import instance_rgb


def load_config(path):
    path = Path(path).resolve()
    config = json.loads(path.read_text(encoding="utf-8-sig"))
    root = Path(config["global_root"]) / config["project_directory"]
    if not root.is_absolute():
        root = path.parent / root
    config["resolved_paths"] = {k: str(root / v) for k, v in config["paths"].items()}
    return config


def read_image(path):
    path = Path(path)
    if path.suffix.lower() in (".tif", ".tiff"):
        a = tifffile.imread(path)
    else:
        with Image.open(path) as im:
            a = np.asarray(im).copy()
    if not (a.ndim == 2 or (a.ndim == 3 and a.shape[-1] in (3, 4))):
        raise ValueError("This demo accepts single 2D grayscale or RGB images, not stacks.")
    if a.ndim == 3 and a.shape[-1] == 4:
        a = a[..., :3]
    if not np.isfinite(a).all():
        raise ValueError("Image contains NaN or infinite values.")
    return a


def normalized_image(a):
    if np.issubdtype(a.dtype, np.unsignedinteger):
        return a.astype(np.float32) / np.iinfo(a.dtype).max
    f = a.astype(np.float32)
    if f.min() < 0 or f.max() > 1:
        lo, hi = float(f.min()), float(f.max())
        f = (f - lo) / (hi - lo) if hi > lo else np.zeros_like(f)
    return f


def illumination_radius(image):
    """Resolution-based preparation used by both single-image and project detection."""
    return int(image.shape[0]) * int(image.shape[1]) // 10000


def illumination_correct(image, radius):
    """Notebook formula with Wolfram's default discrete Bessel Gaussian kernel.

    Wolfram radius r corresponds to sigma r/2, with fixed edge padding.
    The notebook chooses floor(width*height/10000) as its radius.
    """
    a = normalized_image(image)
    radius = int(radius)
    if radius <= 0:
        return a.copy()
    import cv2
    # exp(-sigma**2) * BesselI[k, sigma**2], normalized after truncation.
    # ive evaluates the exponentially scaled Bessel function without overflow.
    kernel = ive(np.arange(-radius, radius + 1), (radius / 2)**2)
    kernel = (kernel / kernel.sum()).astype(np.float32)
    bg = cv2.sepFilter2D(a, -1, kernel, kernel, borderType=cv2.BORDER_REPLICATE)
    return np.clip(a - bg + bg.mean(), 0, 1)


def model_rgb(image):
    a = np.round(normalized_image(image) * 255).astype(np.uint8)
    if a.ndim == 2:
        a = np.repeat(a[..., None], 3, axis=-1)
    return np.ascontiguousarray(a)


def filter_boxes(boxes, scores, confidence=.25, containment=.9, iou=.5):
    """Original containment -> confidence -> NMS order; handle empty/tied boxes."""
    import torch
    from torchvision.ops import nms

    boxes = np.asarray(boxes, dtype=np.float32).reshape(-1, 4)
    scores = np.asarray(scores, dtype=np.float32).reshape(-1)
    if len(boxes) != len(scores):
        raise ValueError("Boxes and scores must have matching lengths.")
    valid = np.isfinite(boxes).all(1) & np.isfinite(scores)
    valid &= (boxes[:, 2:] > boxes[:, :2]).all(1)
    boxes, scores = boxes[valid], scores[valid]
    counts = {"candidates": len(boxes)}
    area = np.prod(boxes[:, 2:] - boxes[:, :2], axis=1)
    keep = np.ones(len(boxes), dtype=bool)
    # Each pair once prevents equal-area duplicates deleting each other.
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            intersection = np.prod(np.maximum(0, np.minimum(boxes[i, 2:], boxes[j, 2:])
                                              - np.maximum(boxes[i, :2], boxes[j, :2])))
            if intersection / (min(area[i], area[j]) + 1e-6) > containment:
                loser = i if area[i] < area[j] else j
                if area[i] == area[j]:
                    loser = i if scores[i] < scores[j] else j
                keep[loser] = False
    boxes, scores = boxes[keep], scores[keep]
    counts["after_containment"] = len(boxes)
    keep = scores > confidence
    boxes, scores = boxes[keep], scores[keep]
    counts["after_confidence"] = len(boxes)
    keep = nms(torch.from_numpy(boxes), torch.from_numpy(scores), iou).numpy()
    counts["after_nms"] = len(keep)
    return boxes[keep], scores[keep], counts


def boxes_to_shapes(boxes):
    return [np.array([[y1, x1], [y1, x2], [y2, x2], [y2, x1]])
            for x1, y1, x2, y2 in boxes]


def shapes_to_boxes(shapes, image_shape):
    h, w = image_shape[:2]
    boxes = []
    for index, shape in enumerate(shapes):
        a = np.asarray(shape)
        if a.shape != (4, 2) or not np.isfinite(a).all():
            raise ValueError("Please use 2D rectangles for all box prompts.")
        # Napari uses float32 vertices. Resize/move transforms can leave tiny
        # differences between nominally equal coordinates. Counting rounded
        # unique values falsely rejects these at rounding boundaries.
        low, high = a.min(0), a.max(0)
        upper = abs(a - high) < abs(a - low)
        snapped = np.where(upper, high, low)
        if np.max(abs(a - snapped)) > 0.01:
            raise ValueError(f"Box {index + 1} is not rectangular. Delete and redraw this box.")
        y1, x1 = np.maximum(a.min(0), [0, 0])
        y2, x2 = np.minimum(a.max(0), [h, w])
        if x2 - x1 < 1 or y2 - y1 < 1:
            raise ValueError("A box is empty or outside the image. Delete or resize it.")
        if len(np.unique(upper, axis=0)) != 4:
            raise ValueError(f"Box {index + 1} has invalid corners. Delete and redraw this box.")
        boxes.append([x1, y1, x2, y2])
    return np.asarray(boxes, dtype=np.float32).reshape(-1, 4)


class Engine:
    def __init__(self, config):
        self.config = config
        self.predictor = None
        self.yolo = None
        self.yolo_path = None

    def device(self):
        import torch
        requested = self.config["inference"]["device"]
        if requested not in {"auto", "cpu", "cuda"}:
            raise ValueError("Device must be auto, cpu, or cuda.")
        if requested == "cuda" and not torch.cuda.is_available():
            raise RuntimeError("CUDA is unavailable. Install CUDA-enabled PyTorch or set device to cpu in workflow_config.json.")
        return "cuda" if requested == "auto" and torch.cuda.is_available() else ("cpu" if requested == "auto" else requested)

    def release_sam(self):
        import torch
        self.predictor = None
        gc.collect()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    def release_yolo(self):
        self.yolo = None
        self.yolo_path = None
        gc.collect()

    def detect(self, image, keep_model=False):
        import torch
        from ultralytics import YOLO
        if self.predictor is not None:
            self.release_sam()
        cfg = self.config["inference"]
        path = self.config["resolved_paths"]["yolo_model"]
        if self.yolo is None or self.yolo_path != path:
            self.yolo = YOLO(path)
            self.yolo_path = path
        model = self.yolo
        try:
            # Ultralytics ndarray input is BGR; SAM input is RGB.
            result = model.predict(model_rgb(image)[..., ::-1].copy(),
                                   conf=cfg["yolo_candidate_confidence"],
                                   imgsz=cfg["yolo_imgsz"], device=self.device(),
                                   verbose=False)[0]
            boxes = result.boxes.xyxy.cpu().numpy().copy()
            scores = result.boxes.conf.cpu().numpy().copy()
            return filter_boxes(boxes, scores, cfg["confidence_threshold"],
                                cfg["containment_threshold"], cfg["nms_iou"])
        finally:
            del model
            if not keep_model:
                self.release_yolo()
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()

    def segment(self, image, boxes, scores, progress=lambda message: None):
        import torch
        from segment_anything import SamPredictor, sam_model_registry
        cfg = self.config["inference"]
        device = self.device()
        boxes = np.asarray(boxes, np.float32).reshape(-1, 4)
        scores = np.asarray(scores, np.float32)
        if len(boxes) != len(scores) or not np.isfinite(scores).all():
            raise ValueError("Invalid box confidences.")
        instance = np.zeros(image.shape[:2], dtype=np.uint32)
        score_map = np.full(image.shape[:2], -np.inf, dtype=np.float32)
        records = []
        if not len(boxes):
            return instance, records
        try:
            if self.predictor is None:
                progress("Loading SAM ViT-H on " + device + "…")
                model = sam_model_registry[cfg["sam_type"]](checkpoint=self.config["resolved_paths"]["sam_model"])
                model.to(device=device)
                self.predictor = SamPredictor(model)
            with torch.inference_mode():
                progress("Encoding image with SAM…")
                self.predictor.set_image(model_rgb(image))
                batch = max(1, int(cfg["sam_batch_size"]))
                for start in range(0, len(boxes), batch):
                    stop = min(start + batch, len(boxes))
                    prompts = torch.as_tensor(boxes[start:stop], device=device)
                    transformed = self.predictor.transform.apply_boxes_torch(prompts, image.shape[:2])
                    masks, quality, _ = self.predictor.predict_torch(
                        point_coords=None, point_labels=None, boxes=transformed, multimask_output=True)
                    best = quality.argmax(dim=1)
                    for local, idx in enumerate(range(start, stop)):
                        mask = masks[local, best[local]].cpu().numpy()
                        sam_score = float(quality[local, best[local]].clamp(0, 1))
                        combined = cfg["yolo_weight"] * float(scores[idx]) + (1-cfg["yolo_weight"]) * sam_score
                        update = mask & (combined > score_map)
                        instance[update] = idx + 1
                        score_map[update] = combined
                        records.append({"label_id": idx + 1, "box_confidence": float(scores[idx]),
                                        "sam_score": sam_score, "combined_score": combined})
                    del masks, quality, prompts, transformed
                    progress(f"Segmented {stop}/{len(boxes)} boxes")
            return instance, records
        except torch.cuda.OutOfMemoryError as exc:
            self.release_sam()
            raise RuntimeError("GPU memory ran out. Close other GPU applications and retry; use device=cpu in the config if necessary.") from exc


def export_result(labels, output_root, image_path, pixel_size=1., unit="px", metadata=None, formats=None, folder_name=None, overwrite=False):
    labels = np.asarray(labels)
    if labels.ndim != 2 or not np.issubdtype(labels.dtype, np.integer) or labels.min() < 0:
        raise ValueError("Expected a nonnegative 2D integer label image.")
    if not np.isfinite(pixel_size) or pixel_size <= 0 or not unit.strip():
        raise ValueError("Provide a positive pixel size and a measurement unit.")
    formats = set(("labels", "binary", "csv") if formats is None else formats)
    if not formats or formats - {"labels", "binary", "colored", "csv", "xlsx"}:
        raise ValueError("Choose at least one supported export format.")
    folder_name = Path(image_path).stem if folder_name is None else folder_name
    if (not folder_name or folder_name in (".", "..") or any(c in folder_name for c in '<>:"/\\|?*')
            or folder_name.endswith((".", " "))):
        raise ValueError("Enter a valid export subfolder name, without path separators.")
    directory = Path(output_root) / folder_name
    if directory.is_symlink() or (directory.exists() and not directory.is_dir()):
        raise ValueError("Export destination must be a regular folder.")
    if directory.exists() and not overwrite:
        raise FileExistsError(f"Export folder already exists: {directory}")
    columns = ["label_id", "area_px2", "perimeter_px", "centroid_x_px", "centroid_y_px",
               "area", "perimeter", "centroid_x", "centroid_y", "unit"]
    rows = []
    if formats & {"csv", "xlsx"}:
        for prop in regionprops(labels):
            y, x = prop.centroid
            rows.append([int(prop.label), float(prop.area), float(prop.perimeter), float(x), float(y),
                prop.area * pixel_size**2, prop.perimeter * pixel_size,
                x * pixel_size, y * pixel_size, unit])
    payload = dict(metadata or {})
    payload.update(image_path=str(image_path), pixel_size=pixel_size, unit=unit,
                   export_formats=sorted(formats), output_root=str(Path(output_root).resolve()),
                   colored_mask="RGB palette on black background; no selection highlight; use labels.tif for IDs",
                   measurement_source="final edited label pixels", centroid_origin="zero-based top-left; x=column, y=row",
                   perimeter_method="skimage.measure.regionprops perimeter, 4-connectivity")
    # Publish only a complete export; a failed writer leaves no partial result.
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".export_", dir=root) as temporary:
        stage = Path(temporary) / "result"
        stage.mkdir()
        if "labels" in formats:
            tifffile.imwrite(stage / "labels.tif", labels.astype(np.uint32), photometric="minisblack", metadata={"axes": "YX"})
        if "binary" in formats:
            tifffile.imwrite(stage / "binary.tif", (labels > 0).astype(np.uint8) * 255, photometric="minisblack", metadata={"axes": "YX"})
        if "colored" in formats:
            tifffile.imwrite(stage / "colored_instances.tif", instance_rgb(labels), photometric="rgb", metadata={"axes": "YXS"})
        if "csv" in formats:
            with (stage / "measurements.csv").open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(columns)
                writer.writerows(rows)
        if "xlsx" in formats:
            write_measurement_xlsx(stage / "measurements.xlsx", columns, rows)
        (stage / "session.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
        if not directory.exists():
            stage.rename(directory)
        else:
            managed = {"labels.tif", "binary.tif", "colored_instances.tif", "measurements.csv", "measurements.xlsx", "session.json"}
            backup = Path(temporary) / "backup"
            backup.mkdir()
            written = []
            try:
                for name in managed:
                    old = directory / name
                    if old.exists():
                        if not old.is_file() or old.is_symlink():
                            raise ValueError(f"Cannot replace non-regular output: {old}")
                        old.rename(backup / name)
                for new in stage.iterdir():
                    new.rename(directory / new.name)
                    written.append(directory / new.name)
            except Exception:
                for new in written:
                    new.unlink()
                for old in backup.iterdir():
                    old.rename(directory / old.name)
                raise
    return directory


def write_combined_measurements(root, columns, rows, overwrite):
    """Stage both combined files and roll back either publication on failure."""
    names = ('all_measurements.csv', 'all_measurements.xlsx')
    for name in names:
        target = root / name
        if target.is_symlink() or (target.exists() and not target.is_file()):
            raise ValueError(f'Cannot replace non-regular output: {target}')
        if target.exists() and not overwrite:
            raise FileExistsError('Combined files already exist. Choose Replace existing results to update them.')
    with tempfile.TemporaryDirectory(prefix='.merge_', dir=root) as temporary:
        stage = Path(temporary)
        with (stage / names[0]).open('w', newline='', encoding='utf-8') as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            writer.writerows(rows)
        write_measurement_xlsx(stage / names[1], columns, rows)
        backup = stage / 'backup'
        backup.mkdir()
        written = []
        try:
            for name in names:
                if (root / name).exists():
                    (root / name).rename(backup / name)
            for name in names:
                (stage / name).rename(root / name)
                written.append(root / name)
        except Exception:
            for target in written:
                target.unlink()
            for old in backup.iterdir():
                old.rename(root / old.name)
            raise


def write_measurement_xlsx(path, columns, rows):
    """Application runtime export; numeric cells remain numbers in Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
        # Names and units are text, never spreadsheet formulas.
        for index, value in enumerate(row, 1):
            if isinstance(value, str):
                sheet.cell(sheet.max_row, index).data_type = 's'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="376B72")
        cell.alignment = Alignment(horizontal="left")
    for index, heading in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(heading) + 3)
    for row in sheet.iter_rows(min_row=2):
        for heading, cell in zip(columns, row):
            if heading != 'label_id' and isinstance(cell.value, (int, float)):
                cell.number_format = "0.000"
    workbook.save(path)
    workbook.close()
